
import sys, sqlite3, csv, os
from datetime import date, datetime
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
    QLabel, QPushButton, QLineEdit, QComboBox, QDoubleSpinBox, QDateEdit,
    QTableWidget, QTableWidgetItem, QMessageBox, QFileDialog, QDialog,
    QFormLayout, QSpinBox, QHeaderView, QGroupBox
)

APP_DIR = Path.home() / "DagOtoAsistan"
APP_DIR.mkdir(exist_ok=True)
DB = APP_DIR / "dag_oto.db"

def money(v):
    return f"{v:,.2f} ₺".replace(",", "X").replace(".", ",").replace("X", ".")

class DBManager:
    def __init__(self):
        self.con = sqlite3.connect(DB)
        self.con.execute("PRAGMA foreign_keys=ON")
        self.init()

    def init(self):
        self.con.executescript("""
        CREATE TABLE IF NOT EXISTS customers(
            id INTEGER PRIMARY KEY, name TEXT NOT NULL, phone TEXT DEFAULT '',
            note TEXT DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS vehicles(
            id INTEGER PRIMARY KEY, customer_id INTEGER, plate TEXT UNIQUE NOT NULL,
            brand TEXT DEFAULT 'Ford', model TEXT DEFAULT '', vin TEXT DEFAULT '',
            km INTEGER DEFAULT 0,
            FOREIGN KEY(customer_id) REFERENCES customers(id)
        );
        CREATE TABLE IF NOT EXISTS parts(
            id INTEGER PRIMARY KEY, code TEXT UNIQUE NOT NULL, name TEXT NOT NULL,
            price REAL DEFAULT 0, stock REAL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS jobs(
            id INTEGER PRIMARY KEY, job_date TEXT NOT NULL, customer_id INTEGER,
            vehicle_id INTEGER, description TEXT DEFAULT '', labor REAL DEFAULT 0,
            total REAL DEFAULT 0, paid REAL DEFAULT 0, note TEXT DEFAULT '',
            FOREIGN KEY(customer_id) REFERENCES customers(id),
            FOREIGN KEY(vehicle_id) REFERENCES vehicles(id)
        );
        CREATE TABLE IF NOT EXISTS job_parts(
            id INTEGER PRIMARY KEY, job_id INTEGER, part_id INTEGER,
            qty REAL DEFAULT 1, unit_price REAL DEFAULT 0,
            FOREIGN KEY(job_id) REFERENCES jobs(id) ON DELETE CASCADE,
            FOREIGN KEY(part_id) REFERENCES parts(id)
        );
        CREATE TABLE IF NOT EXISTS payments(
            id INTEGER PRIMARY KEY, customer_id INTEGER, payment_date TEXT NOT NULL,
            amount REAL NOT NULL, job_id INTEGER, note TEXT DEFAULT '',
            FOREIGN KEY(customer_id) REFERENCES customers(id),
            FOREIGN KEY(job_id) REFERENCES jobs(id)
        );
        CREATE TABLE IF NOT EXISTS reminders(
            id INTEGER PRIMARY KEY, customer_id INTEGER, due_date TEXT,
            note TEXT, done INTEGER DEFAULT 0,
            FOREIGN KEY(customer_id) REFERENCES customers(id)
        );
        """)
        self.con.commit()

    def customer_id(self, name):
        row = self.con.execute("SELECT id FROM customers WHERE name=?", (name.strip(),)).fetchone()
        if row: return row[0]
        cur = self.con.execute("INSERT INTO customers(name) VALUES(?)", (name.strip(),))
        self.con.commit()
        return cur.lastrowid

    def vehicle_id(self, plate, customer_id, model="", vin="", km=0):
        plate = plate.strip().upper()
        row = self.con.execute("SELECT id FROM vehicles WHERE plate=?", (plate,)).fetchone()
        if row:
            self.con.execute("UPDATE vehicles SET customer_id=?, model=?, vin=?, km=? WHERE id=?",
                             (customer_id, model, vin, km, row[0]))
            self.con.commit()
            return row[0]
        cur = self.con.execute(
            "INSERT INTO vehicles(customer_id,plate,model,vin,km) VALUES(?,?,?,?,?)",
            (customer_id, plate, model, vin, km))
        self.con.commit()
        return cur.lastrowid

    def add_job(self, d, customer, plate, model, vin, km, desc, labor, total, paid, part_code="", part_name="", qty=1, part_price=0):
        cid = self.customer_id(customer)
        vid = self.vehicle_id(plate, cid, model, vin, km)
        cur = self.con.execute(
            "INSERT INTO jobs(job_date,customer_id,vehicle_id,description,labor,total,paid) VALUES(?,?,?,?,?,?,?)",
            (d, cid, vid, desc, labor, total, paid))
        jid = cur.lastrowid
        if part_code:
            row = self.con.execute("SELECT id,name,price FROM parts WHERE code=?", (part_code.strip(),)).fetchone()
            if row:
                pid = row[0]
                up = part_price or row[2]
                self.con.execute("INSERT INTO job_parts(job_id,part_id,qty,unit_price) VALUES(?,?,?,?)",
                                 (jid, pid, qty, up))
                self.con.execute("UPDATE parts SET stock=stock-? WHERE id=?", (qty, pid))
            else:
                pid = self.con.execute("INSERT INTO parts(code,name,price,stock) VALUES(?,?,?,0)",
                                       (part_code.strip(), part_name or "Yeni parça", part_price)).lastrowid
                self.con.execute("INSERT INTO job_parts(job_id,part_id,qty,unit_price) VALUES(?,?,?,?)",
                                 (jid, pid, qty, part_price))
        if paid:
            self.con.execute("INSERT INTO payments(customer_id,payment_date,amount,job_id,note) VALUES(?,?,?,?,?)",
                             (cid, d, paid, jid, "Servis kaydı sırasında alınan ödeme"))
        self.con.commit()
        return jid

    def customer_balance(self, cid):
        billed = self.con.execute("SELECT COALESCE(SUM(total),0) FROM jobs WHERE customer_id=?", (cid,)).fetchone()[0]
        paid = self.con.execute("SELECT COALESCE(SUM(amount),0) FROM payments WHERE customer_id=?", (cid,)).fetchone()[0]
        return billed, paid, billed-paid

db = DBManager()

class ServiceDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Yeni Servis Kaydı")
        self.resize(620, 520)
        f = QFormLayout(self)
        self.dt = QDateEdit(); self.dt.setCalendarPopup(True); self.dt.setDate(datetime.now().date())
        self.customer = QLineEdit()
        self.plate = QLineEdit(); self.plate.setPlaceholderText("Örn. 44 ABC 123")
        self.model = QLineEdit(); self.model.setPlaceholderText("Örn. F-MAX / 4145 / 4142")
        self.vin = QLineEdit()
        self.km = QSpinBox(); self.km.setMaximum(9999999)
        self.desc = QLineEdit()
        self.labor = QDoubleSpinBox(); self.labor.setMaximum(999999999); self.labor.setDecimals(2)
        self.total = QDoubleSpinBox(); self.total.setMaximum(999999999); self.total.setDecimals(2)
        self.paid = QDoubleSpinBox(); self.paid.setMaximum(999999999); self.paid.setDecimals(2)
        self.part_code = QLineEdit(); self.part_code.setPlaceholderText("Parça kodu")
        self.part_name = QLineEdit(); self.part_name.setPlaceholderText("Kod bulunamazsa parça adı")
        self.qty = QDoubleSpinBox(); self.qty.setMinimum(0.01); self.qty.setMaximum(99999); self.qty.setValue(1)
        self.part_price = QDoubleSpinBox(); self.part_price.setMaximum(999999999); self.part_price.setDecimals(2)

        f.addRow("Tarih", self.dt); f.addRow("Müşteri", self.customer); f.addRow("Plaka *", self.plate)
        f.addRow("Araç / Model", self.model); f.addRow("Şasi / VIN", self.vin); f.addRow("Kilometre", self.km)
        f.addRow("Yapılan iş", self.desc); f.addRow("İşçilik", self.labor); f.addRow("Toplam", self.total)
        f.addRow("Bu işlemde alınan ödeme", self.paid)
        box = QGroupBox("Parça (isteğe bağlı)")
        pf = QFormLayout(box); pf.addRow("Parça kodu", self.part_code); pf.addRow("Parça adı", self.part_name)
        pf.addRow("Adet", self.qty); pf.addRow("Birim fiyat", self.part_price); f.addRow(box)
        self.save = QPushButton("Kaydet"); self.save.clicked.connect(self.save_job); f.addRow(self.save)
        self.part_code.editingFinished.connect(self.lookup_part)

    def lookup_part(self):
        code = self.part_code.text().strip()
        if not code: return
        row = db.con.execute("SELECT name,price FROM parts WHERE code=?", (code,)).fetchone()
        if row:
            self.part_name.setText(row[0]); self.part_price.setValue(row[1])
        else:
            self.part_name.setPlaceholderText("Bu kod katalogda yok — kaydederken yeni parça olarak eklenir")

    def save_job(self):
        if not self.customer.text().strip() or not self.plate.text().strip():
            QMessageBox.warning(self, "Eksik bilgi", "Müşteri ve plaka zorunludur.")
            return
        if self.paid.value() > self.total.value():
            QMessageBox.warning(self, "Hatalı ödeme", "Ödeme toplam tutardan büyük olamaz.")
            return
        db.add_job(self.dt.date().toString("yyyy-MM-dd"), self.customer.text(), self.plate.text(),
                   self.model.text(), self.vin.text(), self.km.value(), self.desc.text(),
                   self.labor.value(), self.total.value(), self.paid.value(),
                   self.part_code.text(), self.part_name.text(), self.qty.value(), self.part_price.value())
        self.accept()

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Dağ Oto Muhasebe & Servis")
        self.resize(1250, 760)
        self.build()
        self.refresh()

    def build(self):
        central = QWidget(); self.setCentralWidget(central)
        outer = QHBoxLayout(central)
        side = QVBoxLayout()
        brand = QLabel("DAĞ OTO\nMUHASEBE & SERVİS"); brand.setObjectName("brand")
        side.addWidget(brand)
        for text, slot in [
            ("🏠 Ana Sayfa", self.refresh), ("📅 Günlük İşler", self.refresh),
            ("👥 Müşteriler", self.show_customers), ("🚛 Araçlar / Plakalar", self.show_vehicles),
            ("🔧 Yeni Servis", self.new_service), ("📦 Parçalar", self.show_parts),
            ("💳 Ödemeler", self.show_payments), ("📊 Raporlar", self.refresh),
            ("📥 Excel Aktar", self.import_excel), ("💾 Yedekle", self.backup)
        ]:
            b=QPushButton(text); b.clicked.connect(slot); side.addWidget(b)
        side.addStretch()
        outer.addLayout(side, 0)

        self.body = QVBoxLayout(); outer.addLayout(self.body, 1)
        self.header = QLabel(); self.header.setObjectName("header"); self.body.addWidget(self.header)
        self.cards = QHBoxLayout(); self.body.addLayout(self.cards)
        self.table = QTableWidget(); self.body.addWidget(self.table, 1)

    def clear_body(self):
        while self.cards.count():
            item=self.cards.takeAt(0)
            if item.widget(): item.widget().deleteLater()

    def card(self, title, value):
        w=QGroupBox(); l=QVBoxLayout(w); a=QLabel(title); b=QLabel(value); b.setObjectName("big")
        l.addWidget(a); l.addWidget(b); self.cards.addWidget(w)

    def refresh(self):
        self.clear_body()
        today=date.today().isoformat()
        self.header.setText(f"Bugün — {today}")
        billed, paid, balance = self.con_totals()
        today_jobs=self.con_scalar("SELECT COUNT(*) FROM jobs WHERE job_date=?", (today,))
        self.card("Bugünkü İşler", str(today_jobs))
        self.card("Toplam Ciro", money(billed))
        self.card("Toplam Ödeme", money(paid))
        self.card("Açık Borç", money(balance))
        rows=db.con.execute("""
            SELECT j.job_date,c.name,v.plate,v.model,j.description,j.total,
                   j.total-COALESCE((SELECT SUM(p.amount) FROM payments p WHERE p.job_id=j.id),0)
            FROM jobs j LEFT JOIN customers c ON c.id=j.customer_id
            LEFT JOIN vehicles v ON v.id=j.vehicle_id ORDER BY j.job_date DESC,j.id DESC LIMIT 100
        """).fetchall()
        headers=["Tarih","Müşteri","Plaka","Araç","Yapılan İş","Tutar","Kalan"]
        self.fill_table(headers, rows)

    def con_scalar(self, q, args=()):
        return db.con.execute(q,args).fetchone()[0]

    def con_totals(self):
        billed=self.con_scalar("SELECT COALESCE(SUM(total),0) FROM jobs")
        paid=self.con_scalar("SELECT COALESCE(SUM(amount),0) FROM payments")
        return billed,paid,billed-paid

    def fill_table(self, headers, rows):
        self.table.clear(); self.table.setColumnCount(len(headers)); self.table.setHorizontalHeaderLabels(headers)
        self.table.setRowCount(len(rows))
        for r,row in enumerate(rows):
            for c,val in enumerate(row):
                self.table.setItem(r,c,QTableWidgetItem("" if val is None else (money(val) if isinstance(val,(float,int)) and c in [5,6] else str(val))))
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setStretchLastSection(True)

    def new_service(self):
        if ServiceDialog(self).exec(): self.refresh()

    def show_customers(self):
        rows=db.con.execute("SELECT id,name,phone FROM customers ORDER BY name").fetchall()
        out=[]
        for cid,name,phone in rows:
            billed,paid,balance=db.customer_balance(cid); out.append((name,phone,billed,paid,balance))
        self.header.setText("Müşteriler / Cari Hesaplar")
        self.fill_table(["Müşteri","Telefon","Toplam Borç","Ödenen","Kalan"],out)

    def show_vehicles(self):
        rows=db.con.execute("""SELECT v.plate,c.name,v.model,v.vin,v.km,
             COUNT(j.id),COALESCE(SUM(j.total),0)
             FROM vehicles v LEFT JOIN customers c ON c.id=v.customer_id
             LEFT JOIN jobs j ON j.vehicle_id=v.id GROUP BY v.id ORDER BY v.plate""").fetchall()
        self.header.setText("Araçlar / Plakalar")
        self.fill_table(["Plaka","Müşteri","Araç","VIN","KM","Servis Sayısı","Toplam"],rows)

    def show_parts(self):
        rows=db.con.execute("SELECT code,name,price,stock FROM parts ORDER BY code").fetchall()
        self.header.setText("Parça Kataloğu / Stok")
        self.fill_table(["Parça Kodu","Parça Adı","Birim Fiyat","Stok"],rows)

    def show_payments(self):
        rows=db.con.execute("""SELECT p.payment_date,c.name,p.amount,p.note
             FROM payments p JOIN customers c ON c.id=p.customer_id ORDER BY p.payment_date DESC,p.id DESC""").fetchall()
        self.header.setText("Ödemeler")
        self.fill_table(["Tarih","Müşteri","Ödeme","Not"],rows)

    def import_excel(self):
        path,_=QFileDialog.getOpenFileName(self,"Excel/CSV seç","","Excel (*.xlsx);;CSV (*.csv)")
        if not path:return
        if path.lower().endswith(".xlsx"):
            try:
                from openpyxl import load_workbook
                ws=load_workbook(path, data_only=True).active
                rows=list(ws.iter_rows(values_only=True))
                if not rows: return
                headers=[str(x or "").strip().lower() for x in rows[0]]
                aliases={
                    "tarih":["tarih","date"],"müşteri":["müşteri","musteri","cari"],
                    "plaka":["plaka","plate"],"araç":["araç","arac","model"],
                    "işlem":["işlem","islem","yapılan iş"],"tutar":["tutar","toplam"],
                    "ödeme":["ödeme","odeme","ödenen"],"parça kodu":["parça kodu","parca kodu","kod"]
                }
                def find(key):
                    for a in aliases[key]:
                        if a in headers:return headers.index(a)
                    return None
                idx={k:find(k) for k in aliases}
                imported=0; missing=[]
                for n,row in enumerate(rows[1:],2):
                    def val(k):
                        i=idx[k]; return "" if i is None or i>=len(row) or row[i] is None else str(row[i]).strip()
                    customer,plate=val("müşteri"),val("plaka")
                    if not customer or not plate:
                        missing.append(n); continue
                    d=val("tarih") or date.today().isoformat()
                    total=float(val("tutar").replace(",",".")) if val("tutar") else 0
                    paid=float(val("ödeme").replace(",",".")) if val("ödeme") else 0
                    db.add_job(d,customer,plate,val("araç"),"","0",val("işlem"),0,total,paid,val("parça kodu"))
                    imported+=1
                QMessageBox.information(self,"Excel aktarımı",f"{imported} kayıt aktarıldı. Eksik müşteri/plaka nedeniyle atlanan satır: {len(missing)}")
                self.refresh()
            except Exception as e:
                QMessageBox.critical(self,"Aktarım hatası",str(e))
        else:
            QMessageBox.information(self,"CSV","CSV aktarımı için ilk satır sütun başlığı, sonraki satırlar kayıt olacak şekilde hazırlanmalıdır.")

    def backup(self):
        path,_=QFileDialog.getSaveFileName(self,"Yedek kaydet",f"dag_oto_yedek_{date.today().isoformat()}.db","Veritabanı (*.db)")
        if path:
            import shutil
            db.con.commit(); shutil.copy2(DB,path)
            QMessageBox.information(self,"Yedekleme","Yedek başarıyla oluşturuldu.")

app=QApplication(sys.argv)
app.setStyleSheet("""
QMainWindow{background:#f4f6f8} QWidget{font-size:13px;color:#17212b}
QLabel#brand{background:#10243e;color:white;font-size:18px;font-weight:800;padding:20px;border-radius:8px}
QLabel#header{font-size:22px;font-weight:800;padding:6px}
QLabel#big{font-size:21px;font-weight:800}
QPushButton{padding:10px 12px;border:1px solid #d8dee5;border-radius:7px;background:white;text-align:left}
QPushButton:hover{background:#eef2f5}
QTableWidget{background:white;border:1px solid #dce2e8;border-radius:8px}
QHeaderView::section{padding:8px;background:#10243e;color:white;border:0}
QGroupBox{background:white;border:1px solid #dce2e8;border-radius:8px;padding:12px}
""")
w=MainWindow(); w.show()
sys.exit(app.exec())
